# WebPilot/executor_agents/document_handler_agent/file_parsers.py

from typing import Dict
import logging

logger = logging.getLogger(__name__)

class PDFParser:
    """PDF 파서"""
    
    @staticmethod
    def parse(file_path: str) -> Dict:
        try:
            import pypdf
            
            text = ""
            metadata = {}
            
            with open(file_path, 'rb') as f:
                reader = pypdf.PdfReader(f)
                
                # 메타데이터
                metadata['page_count'] = len(reader.pages)
                if reader.metadata:
                    metadata['author'] = reader.metadata.get('/Author', '')
                    metadata['title'] = reader.metadata.get('/Title', '')
                
                # 텍스트 추출
                for page in reader.pages:
                    text += page.extract_text() + "\n"
            
            logger.info(f"PDF 파싱 완료: {len(text)}자, {metadata['page_count']}페이지")
            
            return {
                "success": True,
                "full_text": text,
                "metadata": metadata
            }
            
        except Exception as e:
            logger.error(f"PDF 파싱 실패: {e}")
            return {
                "success": False,
                "error": str(e)
            }


class DOCXParser:
    """DOCX 파서"""
    
    @staticmethod
    def parse(file_path: str) -> Dict:
        try:
            import docx
            
            doc = docx.Document(file_path)
            
            # 텍스트 추출
            text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
            
            # 메타데이터
            metadata = {
                "page_count": None,  # docx는 페이지 수 정확히 알기 어려움
                "author": doc.core_properties.author or ""
            }
            
            logger.info(f"DOCX 파싱 완료: {len(text)}자")
            
            return {
                "success": True,
                "full_text": text,
                "metadata": metadata
            }
            
        except Exception as e:
            logger.error(f"DOCX 파싱 실패: {e}")
            return {
                "success": False,
                "error": str(e)
            }


class XLSXParser:
    """XLSX 파서"""
    
    @staticmethod
    def parse(file_path: str) -> Dict:
        try:
            import openpyxl
            
            wb = openpyxl.load_workbook(file_path, data_only=True)
            
            text = ""
            for sheet in wb.worksheets:
                text += f"\n=== {sheet.title} ===\n"
                for row in sheet.iter_rows(values_only=True):
                    row_text = '\t'.join([str(cell) if cell is not None else '' for cell in row])
                    if row_text.strip():
                        text += row_text + "\n"
            
            metadata = {
                "sheet_count": len(wb.worksheets),
                "sheet_names": [sheet.title for sheet in wb.worksheets]
            }
            
            logger.info(f"XLSX 파싱 완료: {len(text)}자, {metadata['sheet_count']}개 시트")
            
            return {
                "success": True,
                "full_text": text,
                "metadata": metadata
            }
            
        except Exception as e:
            logger.error(f"XLSX 파싱 실패: {e}")
            return {
                "success": False,
                "error": str(e)
            }


class TXTParser:
    """TXT 파서"""
    
    @staticmethod
    def parse(file_path: str) -> Dict:
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()
            
            metadata = {
                "encoding": "utf-8"
            }
            
            logger.info(f"TXT 파싱 완료: {len(text)}자")
            
            return {
                "success": True,
                "full_text": text,
                "metadata": metadata
            }
            
        except UnicodeDecodeError:
            # 인코딩 재시도
            try:
                with open(file_path, 'r', encoding='cp949') as f:
                    text = f.read()
                
                return {
                    "success": True,
                    "full_text": text,
                    "metadata": {"encoding": "cp949"}
                }
            except:
                pass
        
        except Exception as e:
            logger.error(f"TXT 파싱 실패: {e}")
            return {
                "success": False,
                "error": str(e)
            }